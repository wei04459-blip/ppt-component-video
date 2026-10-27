"""GUI/Agent component entry: semantic data compiled to native Motion Canvas."""
import io, json, math, shutil
from pathlib import Path
from PIL import Image, ImageDraw, ImageOps
from pptx import Presentation
from pptx.util import Inches, Pt
import video_fonts as fonts
import plain_components as plain
import semantic_scene as semantic
_LOADED_RENDERER=semantic.renderer_fingerprint(Path(__file__).resolve().parent)

def ensure_engine_current():
    if semantic.renderer_fingerprint(Path(__file__).resolve().parent)!=_LOADED_RENDERER:
        raise ValueError("工具已更新，请重启工具后重新预览或导出")


FPS=30
PAGE_SECONDS=4


def object_pages(path):
    import auto_components
    return auto_components.extract_pages(path)


def suggested_layout(items):
    """New imports only: a full-width title followed by a natural reading grid."""
    title=None
    for i,item in enumerate(items):
        value=item.get('plain');text=item.get('text','');x,y,w,h=item['target']
        if value is None and text.strip():
            role='title' if y<.24 and w>.5 and h<.28 and len(text)<=100 else 'body'
            value=dict(kind='text',text=text,role=role);item['plain']=value
        if value and value.get('kind')=='text' and value.get('role')=='title' and title is None:title=i
    rest=[i for i in range(len(items)) if i!=title];count=len(rest)
    cols=1 if count<=2 else 2 if count<=6 else 3 if count<=12 else 4;rows=max(1,math.ceil(count/cols));start=.23 if title is not None else .04
    if title is not None:items[title]['before']=[.04,.04,.92,.15]
    cw=.92/cols;ch=(.94-start)/rows
    for order,i in enumerate(rest):items[i]['before']=[.04+order%cols*cw,start+order//cols*ch,cw-.025,ch-.035]
    return items


def analyse(app,payload):
    p=app.project(payload['project']);state=app.read_project(p.name)
    urls=state.get('pages',[])
    if not urls:raise ValueError('请先导入课件')
    folder=Path(state['pages_dir']).parent
    manifest=folder/'objects.json'
    objects=json.loads(manifest.read_text()) if manifest.exists() else {}
    import auto_components,hashlib
    fingerprint=auto_components.analysis_fingerprint();pptx=list(folder.glob('*.pptx'))
    stale=any(o.get('autoNative',{}).get('extractorFingerprint')!=fingerprint for values in objects.values() for o in values)
    if (not objects or stale or not any(o.get('autoNative') for values in objects.values() for o in values)) and len(pptx)==1:
        objects={str(i+1):v for i,v in enumerate(object_pages(pptx[0]))};app.write_json(manifest,objects)
    results=[];cache=p/'automatic-analysis';cache.mkdir(exist_ok=True)
    for i,u in enumerate(urls):
        items=json.loads(json.dumps(objects.get(str(i+1),[])));issues=[]
        if items and items[0].get('extractionError'):issues=[dict(code='native-extraction',message=items[0]['extractionError'])];items=[]
        if payload.get('ocr'):
            source=source_image(app,p,u);sourceHash=hashlib.sha256(source.tobytes()).hexdigest();cached=cache/f'{i+1}-{fingerprint}-{sourceHash}.json'
            if cached.exists():entry=json.loads(cached.read_text());items=entry['items'];issues=entry['issues']
            else:
                if items and any(o.get('autoNative',{}).get('kind')=='image' for o in items):
                    items,extra=auto_components.expand_text_images(items,p/'local-ocr-cache');issues.extend(extra)
                if not items:
                    try:items,issues=auto_components.from_raster(source)
                    except Exception as exc:issues=[dict(code='local-analysis',message=str(exc))]
                app.write_json(cached,dict(analysisFingerprint=fingerprint,sourceHash=sourceHash,items=items,issues=issues))
        # Suggestions are evidence for AI authoring, never an accepted or auto-rendered design.
        if len(items)>39:items=auto_components.pack_items(items)
        results.append(dict(source=u,number=i+1,objects=suggested_layout(items),recommended=False,automation=dict(status='needs-ai',issues=issues,analysisFingerprint=fingerprint,reviewRequired=True,summary=f'辅助材料：{len(items)}个对象；须由AI读完整课件、核内容并设计独立普通版和成品')))
    return dict(pages=results,notice='这些是本地辅助提取结果，不是成品方案。请交给Codex或WorkBuddy读完整课件并生成原生方案，再导入核对实际预览。')




def rectangle(value):
    if not isinstance(value,list) or len(value)!=4 or any(type(n) not in (int,float) or not math.isfinite(n) for n in value):raise ValueError('组件位置无效')
    x,y,w,h=value
    if min(x,y)<0 or min(w,h)<.005 or x+w>1.001 or y+h>1.001:raise ValueError('请将组件完整放在页面内')
    return value


def validate(app,p,v,require_confirm=True):
    pages=v.get('pages',[]);state=app.read_project(p.name)
    if not isinstance(pages,list) or not 1<=len(pages)<=6:raise ValueError('请选择 1–6 页制作视频')
    if v.get('motion') not in ['smooth','stagger','lift']:raise ValueError('请选择动效')
    for page in pages:
        if page.get('automation',{}).get('status')=='blocked' and not (page.get('scene') and page.get('aiReview',{}).get('contentChecked')):raise ValueError('该页自动分析未完成：'+'；'.join(x.get('message','') for x in page['automation'].get('issues',[])))
        if page.get('source') not in state.get('pages',[]):raise ValueError('源课件已更换，请重新读取组件')
        items=page.get('components',[])
        if not 2<=len(items)<=40:raise ValueError('每页请保留 2–40 个组件，截图页需要先画框并填写完整内容')
        if require_confirm and not page.get('confirmed'):raise ValueError('请逐页确认组件和两版布局')
        for item in items:
            rectangle(item.get('target'));rectangle(item.get('before'))
            if not isinstance(item.get('text',''),str) or len(item.get('text',''))>4000:raise ValueError('组件文字最多 4000 字，请分成多个组件')
            plain.validate(plain.spec(item))
        if not any(plain.has_rebuilt_content(plain.spec(item)) for item in items):
            raise ValueError('普通版须重建原文或表格，不能整页复用成品图块')
    return pages


def source_image(app,p,url):
    path=(app.DATA/url[len('/files/'):]).resolve()
    if not path.is_relative_to(p.resolve()):raise ValueError('源图片路径无效')
    return Image.open(path).convert('RGB')


def crop(source,rect):
    x,y,w,h=rect
    return source.crop((round(x*source.width),round(y*source.height),round((x+w)*source.width),round((y+h)*source.height)))


def simple_tile(item,bitmap,size):
    value=plain.spec(item);plain.validate(value)
    return plain.render(value,bitmap,size)


def prepare(app,p,page):
    source=source_image(app,p,page['source']);ratio=source.width/source.height
    size=(960,round(960/ratio)) if ratio>=960/540 else (round(540*ratio),540)
    size=tuple(max(1,n) for n in size)
    source=source.resize(size,Image.Resampling.LANCZOS)
    prepared=[];before=Image.new('RGB',size,'white')
    for ci,item in enumerate(page['components']):
        bitmap=crop(source,item['target']);x,y,w,h=item['before']
        try: tile=simple_tile(item,bitmap,(max(1,round(w*size[0])),max(1,round(h*size[1]))))
        except ValueError as exc: raise ValueError(f'第 {page.get("number","?")} 页，组件 {ci+1}：{exc}') from exc
        before.paste(tile,(round(x*size[0]),round(y*size[1])))
        prepared.append((item,bitmap,tile))
    return source,before,prepared


def pixel_rect(rect,size):
    x,y,w,h=rect
    left,top=round(x*size[0]),round(y*size[1])
    return [left,top,max(1,round((x+w)*size[0])-left),
            max(1,round((y+h)*size[1])-top)]


def semantic_prepare(app,p,page,assets):
    import semantic_scene as semantic
    source=source_image(app,p,page['source'])
    scene=semantic.compile_page(source,page,assets)
    before=semantic.static_image(scene,assets,False)
    after=semantic.static_image(scene,assets,True)
    return scene,before,after


def semantic_preview_board(scene,before,after):
    # This is a layout guide, deliberately not an imitation of Motion Canvas frames.
    thumb=(480,round(480*before.height/before.width));sheet=Image.new('RGB',(960,thumb[1]+32),'#eeeeed');d=ImageDraw.Draw(sheet)
    sheet.paste(before.resize(thumb),(0,32));sheet.paste(after.resize(thumb),(480,32));d.text((8,8),'BEFORE / fixed parent groups',fill='#222222');d.text((488,8),'AFTER / motion must be reviewed in MP4',fill='#222222')
    return sheet


def preview(app,payload):
    ensure_engine_current()
    import base64,semantic_scene as semantic
    p=app.project(payload['project']);v=payload['video'];pages=validate(app,p,v,False);results=[]
    for page in pages:
        paths=semantic.actual_preview(app,p,page)
        results.append(dict({k:'data:image/png;base64,'+base64.b64encode(path.read_bytes()).decode() for k,path in paths.items()},source=page['source']))
    return dict(pages=results,engine='motion-canvas-native',notice='预览来自实际Motion Canvas帧；最终是原生语义重建，样式与源页可能有差异。')


def export(app,job,payload,p):
    ensure_engine_current()
    import semantic_scene as semantic
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    v=payload['video'];pages=validate(app,p,v);out=p/'exports'/job['id'];out.mkdir(parents=True)
    work,assets=semantic.workspace(app,out);scenes=[];proofs=[]
    prs=Presentation();prs.slide_width=Inches(12);prs.slide_height=Inches(6.75)
    for pi,page in enumerate(pages):
        scene,before,after=semantic_prepare(app,p,page,assets);scenes.append(scene);proofs.append(semantic.audit_content(page,scene))
        before.save(out/f'{pi+1:02d}-普通版.png');after.save(out/f'{pi+1:02d}-语义成品版.png')
        source_image(app,p,page['source']).save(out/f'{pi+1:02d}-源页仅核对.png')
        semantic_preview_board(scene,before,after).save(out/f'{pi+1:02d}-布局核对.png')
        slide=prs.slides.add_slide(prs.slide_layouts[6]);factor=min(12/scene['width'],6.75/scene['height']);offsetX=(12-scene['width']*factor)/2;offsetY=(6.75-scene['height']*factor)/2
        for group in scene['groups']:
            gx,gy,gs=group['before']
            for n in group['nodes']:
                if group.get('nodeBefore'):
                    pose=semantic.node_state(group,n,'before')
                    if pose['opacity']<=0:continue
                    left=Inches(offsetX+pose['x']*factor);top=Inches(offsetY+pose['y']*factor)
                    if n['kind']=='text':
                        # Office may measure CJK glyphs wider than the canvas font metrics.
                        width=min(scene['width']-pose['x'],max(pose['w']*1.15,pose['w']+pose['fontSize']*2))
                        height=min(scene['height']-pose['y'],pose['fontSize']*1.8)
                        shape=slide.shapes.add_textbox(left,top,Inches(width*factor),Inches(height*factor));tf=shape.text_frame;tf.margin_left=tf.margin_top=tf.margin_right=tf.margin_bottom=0;tf.word_wrap=False;tf.text=n['text'];p0=tf.paragraphs[0];p0.font.name='Noto Sans CJK SC';p0.font.size=Pt(pose['fontSize']*factor*72);p0.font.bold=pose['fontWeight']>=700;p0.font.color.rgb=RGBColor.from_string(pose['fill'][1:])
                    elif n['kind']=='image':slide.shapes.add_picture(str(work/n['src']),left,top,Inches(pose['w']*factor),Inches(pose['h']*factor))
                    elif n['kind']=='rect':
                        shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,left,top,Inches(pose['w']*factor),Inches(pose['h']*factor));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(pose['fill'][1:]);shape.line.fill.background()
                        if pose.get('stroke'):shape.line.color.rgb=RGBColor.from_string(pose['stroke'][1:]);shape.line.width=Pt(max(.1,pose['lineWidth']*factor*72))
                    continue
                left=Inches(offsetX+(gx+n['x']*gs)*factor);top=Inches(offsetY+(gy+n['y']*gs)*factor)
                if n['kind']=='text':
                    remaining_w=max(1,group['size'][0]-n['x']);remaining_h=max(1,group['size'][1]-n['y'])
                    shape=slide.shapes.add_textbox(left,top,Inches(remaining_w*gs*factor),Inches(min(n['fontSize']*1.8,remaining_h)*gs*factor));tf=shape.text_frame;tf.margin_left=tf.margin_top=tf.margin_right=tf.margin_bottom=0;tf.word_wrap=False;tf.text=n['text'];p0=tf.paragraphs[0];p0.font.name='Noto Sans CJK SC';p0.font.size=Pt(n['fontSize']*gs*factor*72);p0.font.bold=n.get('fontWeight',400)>=700;p0.font.color.rgb=RGBColor.from_string(n.get('fill','#222222')[1:])
                elif n['kind']=='image':slide.shapes.add_picture(str(work/n['src']),left,top,Inches(n['w']*gs*factor),Inches(n['h']*gs*factor))
                elif n['kind']=='rect':
                    shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,left,top,Inches(n['w']*gs*factor),Inches(n['h']*gs*factor));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(n.get('fill','#ffffff')[1:]);shape.line.fill.background()
                    if n.get('stroke'):shape.line.color.rgb=RGBColor.from_string(n['stroke'][1:]);shape.line.width=Pt(.5)
    app.write_json(out/'content-audit.json',dict(boundary='plan-to-scene only; source-slide completeness requires manual visual review',sourceSlideManualReviewRequired=True,pages=proofs))
    deck=out/'普通排版版.pptx';prs.save(deck);app.write_json(out/'组件方案.json',v);app.write_json(out/'原生场景.json',dict(pages=scenes))
    job['message']='正在使用Motion Canvas制作原生组件重组'
    result=semantic.render(app,job,work,scenes,v.get('music',True),out);ensure_engine_current();result['simplePpt']=app.file_url(deck)
    return result
