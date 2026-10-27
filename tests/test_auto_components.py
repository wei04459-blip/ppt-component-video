import base64,io,tempfile,unittest
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches,Pt
import auto_components as a
import semantic_scene as s

class AutomaticPipelineTests(unittest.TestCase):
 def test_complete_native_content_grouping_and_style(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d);prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9);slide=prs.slides.add_slide(prs.slide_layouts[6])
   for i in range(100):
    sh=slide.shapes.add_textbox(Inches(i%10*1.55),Inches(i//10*.8),Inches(1.5),Inches(.6));sh.text='完整对象'+str(i);sh.text_frame.paragraphs[0].font.size=Pt(12)
   prs.save(p/'source.pptx');raw=a.extract_pages(p/'source.pptx')[0];self.assertEqual(len(raw),100);items=a.pack_items(raw);self.assertLessEqual(len(items),39);texts=[part['text'] for item in items for part in (item['plain']['parts'] if item['plain']['kind']=='group' else [item['plain']]) if part['kind']=='text'];self.assertEqual(len(texts),100);self.assertEqual(set(texts),{'完整对象'+str(i) for i in range(100)});[a.plain.validate(item['plain']) for item in items]
 def test_alpha_icon_and_fill_are_not_black_or_opaque(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d);prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6]);sh=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1));sh.text='标题';sh.text_frame.paragraphs[0].font.size=Pt(24)
   im=Image.new('RGBA',(64,64),(0,0,0,0));im.putpixel((32,32),(0,100,200,255));im.save(p/'icon.png');slide.shapes.add_picture(str(p/'icon.png'),Inches(1),Inches(3),Inches(1),Inches(1));prs.save(p/'source.pptx');items=a.extract_pages(p/'source.pptx')[0];photo=next(x for x in items if x['plain']['kind']=='image');asset=Image.open(io.BytesIO(base64.b64decode(photo['autoNative']['imageData'])));self.assertEqual(asset.getpixel((0,0)),(255,255,255))
 def test_whole_page_image_never_becomes_plain_photo(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d);Image.new('RGB',(1600,900),'blue').save(p/'page.png');prs=Presentation();sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.shapes.add_picture(str(p/'page.png'),0,0,prs.slide_width,prs.slide_height);prs.save(p/'source.pptx');self.assertEqual(a.extract_pages(p/'source.pptx'),[[]])
 def test_raster_ocr_failures_are_explicit(self):
  items,issues=a.from_raster(Image.new('RGB',(1600,900),'white'),dict(available=False,issues=['unavailable']))
  self.assertFalse(items);self.assertEqual(issues[0]['code'],'ocr-unavailable')
if __name__=='__main__':unittest.main()
