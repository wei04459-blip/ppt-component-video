import copy,tempfile,unittest
from pathlib import Path
from PIL import Image,ImageChops
import semantic_scene as s
import video_runtime as runtime
class NodeBeforeTests(unittest.TestCase):
 def scene(self):
  return dict(width=960,height=540,groups=[dict(id='dec',size=[20,20],before=[900,20,1],after=[900,20,1],index=-1,nodes=[dict(kind='rect',x=0,y=0,w=20,h=20,fill='#ffffff')]),dict(id='a',size=[300,200],before=[20,20,1],after=[500,250,1],index=0,nodes=[dict(kind='rect',x=0,y=0,w=300,h=200,fill='#ffffff',afterFill='#cceeff',before=dict(opacity=0)),dict(kind='text',id='a',sourceRef=dict(component='a',path='text'),text='同一短句完整保留',x=10,y=10,fontSize=24,fill='#222222',before=dict(x=40,y=60,fontSize=24,fontWeight=400)),dict(kind='image',src='assets/photo.png',x=10,y=60,w=180,h=120,before=dict(x=40,y=350,w=200,h=100))])])
 def test_states_content_bounds_and_ratio(self):
  with tempfile.TemporaryDirectory() as d:
   assets=Path(d)/'assets';assets.mkdir();Image.new('RGB',(180,120),'red').save(assets/'photo.png');scene=self.scene();s.validate(scene);s.prepare_node_states(scene,assets)
   image=scene['groups'][1]['nodes'][2]['states']['before'];self.assertAlmostEqual(image['w']/image['h'],1.5);self.assertEqual(image['x'],65)
   self.assertEqual(scene['groups'][1]['nodes'][1]['states']['before']['x'],40)
   bad=self.scene();bad['groups'][1]['nodes'][2]['before']['w']=0
   with self.assertRaisesRegex(ValueError,'正数'):s.prepare_node_states(bad,assets)
   page=dict(components=[dict(id='a',plain=dict(kind='text',text='同一短句完整保留'))]);self.assertTrue(s.audit_content(page,scene)['identity'])
   for override in [dict(x=float('nan')),dict(x=-50),dict(fontSize=0),dict(text='替换内容')]:
    bad=self.scene();bad['groups'][1]['nodes'][1]['before']=override
    with self.assertRaises(ValueError):s.prepare_node_states(bad,assets)
 def test_ppt_uses_absolute_before_and_hides_panels(self):
  from unittest.mock import patch
  import uuid,base64,io
  import component_video as cv
  from pptx import Presentation
  previous=runtime.DATA
  with tempfile.TemporaryDirectory() as d:
   try:
    runtime.DATA=Path(d);pid=uuid.uuid4().hex;p=runtime.DATA/pid;p.mkdir();source=p/'source.png';Image.new('RGB',(960,540),'white').save(source);url=runtime.file_url(source)
    runtime.write_json(p/'project.json',dict(id=pid,pages=[url],exports=[]));scene=self.scene();raw=io.BytesIO();Image.new('RGB',(180,120),'red').save(raw,format='PNG');photo=scene['groups'][1]['nodes'][2];photo.pop('src');photo['imageData']=base64.b64encode(raw.getvalue()).decode()
    page=dict(source=url,number=1,confirmed=True,scene=scene,components=[dict(id='a',before=[0,0,.5,.5],target=[0,0,.5,.5],plain=dict(kind='text',text='同一短句完整保留'))]);page['components'].append(dict(id='b',before=[.5,.5,.3,.3],target=[.5,.5,.3,.3],plain=dict(kind='text',text='第二内容')))
    scene['groups'][0]['nodes']=[dict(kind='text',id='b',sourceRef=dict(component='b',path='text'),text='第二内容',x=0,y=0,fontSize=20)];scene['groups'][0]['size']=[100,30];scene['groups'][0]['before']=[800,20,1];scene['groups'][0]['after']=[800,20,1]
    job=dict(id=uuid.uuid4().hex,cancel=False,log=str(p/'log'))
    with patch('semantic_scene.render',return_value={}):cv.export(runtime,job,dict(video=dict(mode='components',motion='stagger',music=False,pages=[page])),p)
    deck=Presentation(p/'exports'/job['id']/'普通排版版.pptx');texts=[sh for sh in deck.slides[0].shapes if sh.has_text_frame];node=next(sh for sh in texts if sh.text=='同一短句完整保留');self.assertAlmostEqual(node.left/914400,.5,places=5);self.assertAlmostEqual(node.top/914400,.75,places=5);self.assertEqual(len(deck.slides[0].shapes),3);self.assertGreater(node.width/914400,(len('同一短句完整保留')*24+24)*12/960)
   finally:runtime.DATA=previous
 def test_actual_native_independent_start(self):
  with tempfile.TemporaryDirectory() as d:
   out=Path(d);work,assets=s.workspace(runtime,out);Image.new('RGB',(180,120),'red').save(assets/'photo.png');scene=self.scene();s.validate(scene);s.prepare_node_states(scene,assets)
   frames,audit=s.render_frames(runtime,dict(id='node-before',cancel=False,log=str(out/'render.log')),work,[scene]);self.assertEqual(len(frames),120)
   first=Image.open(frames[0]).convert('RGB');last=Image.open(frames[-1]).convert('RGB');self.assertIsNone(ImageChops.difference(first.crop((62,160,1018,699)),last.crop((62,160,1018,699))).getbbox());self.assertEqual(first.getpixel((140,1220)),(255,0,0));self.assertEqual(last.getpixel((600,1190)),(255,0,0))
if __name__=='__main__':unittest.main()
