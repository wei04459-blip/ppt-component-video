import tempfile,unittest,uuid
from pathlib import Path
from unittest.mock import patch
from PIL import Image
from pptx import Presentation
import component_video as cv
import video_runtime as app

class PptGeometryTests(unittest.TestCase):
 def test_table_text_boxes_and_portrait_text_remain_inside_slide(self):
  previous=app.DATA
  with tempfile.TemporaryDirectory() as folder:
   try:
    app.DATA=Path(folder);pid=uuid.uuid4().hex;p=app.DATA/pid;p.mkdir();sources=[]
    for i,size in enumerate([(1600,1200),(900,1600)]):
     path=p/f'p{i}.png';Image.new('RGB',size,'white').save(path);sources.append(app.file_url(path))
    app.write_json(p/'project.json',dict(id=pid,pages=sources,exports=[]))
    rows=[['id','name','value','notes']]+[[str(i),'Chinese',str(i/10),'full text'] for i in range(1,7)]
    pages=[]
    for i in range(2):
     pages.append(dict(source=sources[i],number=i+1,confirmed=True,components=[dict(id='title',text='',target=[.05,.04,.9,.16],before=[.05,.04,.9,.16],plain=dict(kind='text',text='Test title',role='title')),dict(id='body',text='',target=[.05,.24,.9,.7],before=[.05,.24,.9,.7],plain=dict(kind='table',rows=rows,columns=[1,2,1,3]) if i==0 else dict(kind='text',title='Portrait content',text='Retain text and local group bounds.'))]))
    job=dict(id=uuid.uuid4().hex,cancel=False,log=str(p/'log.txt'))
    with patch('semantic_scene.render',return_value={}):cv.export(app,job,dict(video=dict(mode='components',motion='stagger',music=False,pages=pages)),p)
    deck=Presentation(p/'exports'/job['id']/'\u666e\u901a\u6392\u7248\u7248.pptx');self.assertEqual(len(deck.slides),2)
    for slide in deck.slides:
     for shape in slide.shapes:
      self.assertGreaterEqual(shape.left,-2);self.assertGreaterEqual(shape.top,-2);self.assertLessEqual(shape.left+shape.width,deck.slide_width+2);self.assertLessEqual(shape.top+shape.height,deck.slide_height+2)
   finally:app.DATA=previous

if __name__=='__main__':unittest.main()
