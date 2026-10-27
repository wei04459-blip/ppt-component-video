import copy,json,tempfile,unittest,uuid
from pathlib import Path
from PIL import Image,ImageChops
from pptx import Presentation
from pptx.util import Inches
import test_core
app=test_core.app
import component_video as cv

class ComponentTests(unittest.TestCase):
    def test_native_objects_and_groups(self):
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/'native.pptx';prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
            shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(1));shape.text='真实标题'
            group=slide.shapes.add_group_shape();a=group.shapes.add_textbox(Inches(2),Inches(3),Inches(2),Inches(1));a.text='组合文字';prs.save(p)
            result=cv.object_pages(p)
            self.assertEqual(len(result[0]),2);self.assertEqual(cv.plain.spec(result[0][0])['text'],'真实标题');self.assertEqual(cv.plain.spec(result[0][1])['text'],'组合文字')
            for item in result[0]:cv.rectangle(item['target'])
    def test_validation_preview_and_stale_source(self):
        pid=uuid.uuid4().hex;p=app.DATA/pid;p.mkdir();src=p/'page.png';Image.new('RGB',(1600,900),'#204580').save(src);url=app.file_url(src)
        app.write_json(p/'project.json',{'id':pid,'pages':[url]})
        v={'mode':'components','motion':'stagger','pages':[{'source':url,'confirmed':True,'components':[{'text':'同一份内容','target':[.1,.1,.4,.3],'before':[.05,.1,.8,.6]},{'text':'辅助说明','target':[.55,.1,.35,.2],'before':[.05,.75,.8,.2]}]}]}
        self.assertEqual(len(cv.validate(app,p,v)),1)
        source,before,items=cv.prepare(app,p,v['pages'][0]);self.assertEqual(source.size,(960,540));self.assertIsNotNone(ImageChops.difference(source,before).getbbox())
        bad=copy.deepcopy(v);
        for item in bad['pages'][0]['components']:item['plain']={'kind':'image','assetRole':'photo'}
        with self.assertRaisesRegex(ValueError,'不能整页复用'):cv.validate(app,p,bad)
        bad=copy.deepcopy(v);bad['pages'][0]['confirmed']=False
        with self.assertRaises(ValueError):cv.validate(app,p,bad)
        cv.validate(app,p,bad,False)
        bad=copy.deepcopy(v);bad['pages'][0]['source']='/files/../../etc/passwd'
        with self.assertRaises(ValueError):cv.validate(app,p,bad)
        for r in [[-1,0,.1,.1],[0,0,float('nan'),.1],[.9,.1,.5,.1]]:
            with self.assertRaises(ValueError):cv.rectangle(r)

if __name__=='__main__':unittest.main()
