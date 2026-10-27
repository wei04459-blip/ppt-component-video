import io,unittest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import test_core
import component_video as cv
import plain_components as plain

class PlainContentTests(unittest.TestCase):
    def test_blank_text_cannot_reuse_styled_crop(self):
        with self.assertRaisesRegex(ValueError,'普通版尚未重做'):
            cv.simple_tile({'text':''},Image.new('RGB',(200,100),'blue'),(200,100))
    def test_plain_text_and_table_are_independent_of_source_design(self):
        for value in [{'kind':'text','text':'保持同一份内容 32岁 8人'},
                      {'kind':'table','rows':[['年龄','人数'],['32岁','8人']]}]:
            plain.validate(value)
            a=plain.render(value,Image.new('RGB',(300,160),'blue'),(300,160))
            b=plain.render(value,Image.new('RGB',(300,160),'red'),(300,160))
            self.assertEqual(a.tobytes(),b.tobytes())
            self.assertTrue(all(r==g==b for r,g,b in a.get_flattened_data()))
    def test_table_remains_editable_with_every_cell(self):
        value={'kind':'table','rows':[['姓名','人数'],['XXX','8人']]}
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        plain.add_ppt(slide,value,Image.new('RGB',(300,160)),0,0,Inches(6),Inches(3))
        out=io.BytesIO();prs.save(out);out.seek(0)
        table=Presentation(out).slides[0].shapes[0].table
        self.assertEqual([[c.text for c in row.cells] for row in table.rows],value['rows'])
    def test_photo_must_be_explicit_and_group_bounds_valid(self):
        with self.assertRaises(ValueError):plain.validate({'kind':'image','assetRole':'chart'})
        with self.assertRaises(ValueError):plain.validate({'kind':'group','parts':[{'kind':'text','text':'原文','box':[.9,0,.3,1]}]})
        value={'kind':'image','assetRole':'photo'};plain.validate(value)
        self.assertFalse(plain.has_rebuilt_content(value))
    def test_native_table_import_keeps_rows(self):
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        table=slide.shapes.add_table(2,2,Inches(1),Inches(1),Inches(4),Inches(2)).table
        for i,row in enumerate([['姓名','人数'],['XXX','8人']]):
            for j,cell in enumerate(row):table.cell(i,j).text=cell
        out=io.BytesIO();prs.save(out);out.seek(0)
        self.assertEqual(cv.object_pages(out)[0][0]['plain']['rows'],[['姓名','人数'],['XXX','8人']])
