"""Local automatic content extraction and independent ordinary layout; no slide crops as text."""
import base64,io,math,re
from pathlib import Path
from PIL import Image,ImageDraw
from pptx import Presentation
import video_fonts as fonts
import plain_components as plain

def colour(value,default):
 try:
  if not value or not value.rgb:return default
  raw=str(value.rgb);alpha=next((int(c.get('val'))/100000 for c in value._color._xClr if c.tag.endswith('}alpha')),1)
  return '#'+''.join(f'{round(int(raw[i:i+2],16)*alpha+255*(1-alpha)):02x}' for i in (0,2,4))
 except Exception:return default

def extract_pages(path):
 prs=Presentation(path);pages=[];fingerprint=analysis_fingerprint();W,H=prs.slide_width,prs.slide_height
 for slide in prs.slides:
  content=[];decor=[];issues=[]
  def visit(shapes,ox=0,oy=0,sx=1,sy=1):
   for shape in shapes:
    x=(ox+shape.left*sx)/W;y=(oy+shape.top*sy)/H;w=shape.width*sx/W;h=shape.height*sy/H
    if hasattr(shape,'shapes'):
     xf=shape._element.grpSpPr.xfrm
     if xf is None:issues.append('分组缺少坐标信息');continue
     visit(shape.shapes,ox+shape.left*sx-xf.chOff.x*sx*shape.width/max(1,xf.chExt.cx),oy+shape.top*sy-xf.chOff.y*sy*shape.height/max(1,xf.chExt.cy),sx*shape.width/max(1,xf.chExt.cx),sy*shape.height/max(1,xf.chExt.cy));continue
    if w<=0 or h<=0:continue
    meta=dict(box=[x,y,w,h],id=str(shape.shape_id),name=shape.name)
    if min(x,y)<-.02 or x+w>1.02 or y+h>1.02:
     if getattr(shape,'has_text_frame',False) and shape.text.strip() or getattr(shape,'has_table',False) or getattr(shape,'has_chart',False) or hasattr(shape,'image'):issues.append(f'业务对象{shape.name}超出页面');continue
     x=max(0,x);y=max(0,y);w=min(w,1-x);h=min(h,1-y)
     if min(w,h)<=0:continue
    meta['box']=[max(0,x),max(0,y),min(w,1-max(0,x)),min(h,1-max(0,y))]
    if getattr(shape,'has_table',False):
     table=shape.table;value=dict(kind='table',rows=[[c.text for c in row.cells] for row in table.rows],columns=[col.width for col in table.columns]);meta.update(kind='table',plain=value);content.append(meta)
    elif getattr(shape,'has_chart',False):
     try:
      chart=shape.chart;categories=[c.label for c in chart.plots[0].categories];series=list(chart.series);rows=[['类别']+[s.name for s in series]]+[[str(c)]+[str(s.values[i]) for s in series] for i,c in enumerate(categories)];meta.update(kind='table',plain=dict(kind='table',rows=rows),note='图表以完整数据基础表重建');content.append(meta)
     except Exception:issues.append(f'图表{shape.name}数据无法完整读取')
    elif getattr(shape,'has_text_frame',False) and shape.text.strip():
     paragraphs=list(shape.text_frame.paragraphs);runs=[r for p in paragraphs for r in p.runs];run=next((r for r in runs if r.text.strip()),None);font=run.font if run else paragraphs[0].font
     meta.update(kind='text',plain=dict(kind='text',text=shape.text,role='title' if y<.22 and w>.45 else 'body'),fontSize=(font.size/W*1600 if font.size else (40 if y<.22 else 24)),fontWeight=700 if font.bold else 400,fill=colour(font.color,'#202020'))
     meta['paragraphStyles']=[dict(text=p.text,fontWeight=700 if (next((r.font.bold for r in p.runs if r.font.bold is not None),p.font.bold)) else 400,fill=colour(next((r.font.color for r in p.runs if r.text.strip()),p.font.color),'#303030')) for p in paragraphs]
     try:meta['background']=colour(shape.fill.fore_color,None)
     except Exception:pass
     content.append(meta)
    elif hasattr(shape,'image'):
     try:
      im=Image.open(io.BytesIO(shape.image.blob)).convert('RGBA');white=Image.new('RGBA',im.size,'white');white.alpha_composite(im);im=white.convert('RGB');l,t,r,b=shape.crop_left,shape.crop_top,shape.crop_right,shape.crop_bottom
      im=im.crop((round(l*im.width),round(t*im.height),round((1-r)*im.width),round((1-b)*im.height)));buf=io.BytesIO();im.save(buf,format='PNG');meta.update(kind='image',imageData=base64.b64encode(buf.getvalue()).decode(),plain=dict(kind='image',assetRole='photo'));content.append(meta)
     except Exception:issues.append(f'图片{shape.name}无法读取')
    else:
     try:
      fill=colour(shape.fill.fore_color,None)
      if fill:decor.append(dict(box=meta['box'],fill=fill))
     except Exception:pass
  visit(slide.shapes)
  if len(content)>400:issues.append('超过400个内容对象，请拆页后重试')
  if issues:pages.append([dict(extractionError='；'.join(issues))]);continue
  # A full-slide bitmap has to pass OCR rather than masquerade as a photo.
  raster=not any(o['kind'] in ['text','table'] for o in content) and any(o['kind']=='image' and o['box'][2]*o['box'][3]>.7 for o in content)
  if raster:pages.append([]);continue
  items=[]
  for i,o in enumerate(content):
   items.append(dict(id='auto-'+str(i),name=o['name'],text='',target=o['box'],before=[.04,.04,.92,.92],plain=o['plain'],autoNative={k:v for k,v in o.items() if k not in ['plain','name']}))
  if items:items[0]['autoNative']['decor']=decor
  for item in items:item['autoNative'].update(canvas=[1600,1600*H/W],extractorFingerprint=fingerprint)
  pages.append(items)
 return pages

def split_lines(text,width,fontSize):
 f=fonts.load_font('黑体',round(fontSize));d=ImageDraw.Draw(Image.new('RGB',(1,1)));return plain.wrap(d,text,f,width)

def make_scene(page):
 """Current editable specs are the content truth; original geometry is only a template."""
 items=page['components']
 if any(item['plain']['kind']=='group' and item.get('autoNative',{}).get('members') for item in items):
  leaves=[];mapping={}
  for item in items:
   members=item.get('autoNative',{}).get('members')
   if not members:leaves.append(item);mapping[item['id']]=(item['id'],'');continue
   for i,(part,meta) in enumerate(zip(item['plain']['parts'],members)):
    leaf=dict(id=item['id']+'::'+str(i),name=item['name'],plain={k:v for k,v in part.items() if k!='box'},target=meta['box'],before=item['before'],autoNative=dict(meta))
    mapping[leaf['id']]=(item['id'],f'parts/{i}/');leaf['autoNative']['flowGroup']=item['id'];leaves.append(leaf)
  scene=make_scene(dict(page,components=leaves));merged={};extras=[]
  for g in scene['groups']:
   if g['id'] not in mapping:extras.append(g);continue
   parent,prefix=mapping[g['id']];ox,oy,_=g['after']
   for n in g['nodes']:
    n['x']+=ox;n['y']+=oy
    if n.get('sourceRef'):n['sourceRef']=dict(component=parent,path=prefix+n['sourceRef']['path'])
   merged.setdefault(parent,[]).append(g)
  output=[]
  for parent,gs in merged.items():
   x=min(g['after'][0] for g in gs);y=min(g['after'][1] for g in gs);w=max(g['after'][0]+g['size'][0] for g in gs)-x;h=max(g['after'][1]+g['size'][1] for g in gs)-y;nodes=[n for g in gs for n in g['nodes']]
   for n in nodes:n['x']-=x;n['y']-=y
   output.append(dict(id=parent,size=[w,h],before=[x,y,1],after=[x,y,1],index=len(output),nodes=nodes))
  scene['groups']=extras+output;return scene
 W,H=items[0].get('autoNative',{}).get('canvas',[1600,900]);groups=[];serial=0
 for ci,item in enumerate(items):
  native=item.get('autoNative',{});x,y,w,h=item['target'];x*=W;y*=H;w*=W;h*=H;w=max(4,w);h=max(4,h);value=item['plain'];nodes=[]
  if value['kind']=='text':
   fs=max(14,min(64,native.get('fontSize',28)));lines=split_lines(value['text'],max(20,w-4),fs)
   while (len(lines)*fs*1.22>h or any(fonts.load_font('黑体',round(fs)).getlength(t)>w for t in lines)) and fs>14:
    fs-=1;lines=split_lines(value['text'],max(20,w-4),fs)
   if len(lines)*fs*1.22>h+4:raise ValueError(f'自动成品文字区域太小：{item.get("name",ci)}；请扩大该区域或拆页')
   for li,line in enumerate(lines):
    offset=0
    chunks=[line] if value.get('role')=='title' else split_lines(line,min(w-4,460*fs/26),fs)
    for fragment in chunks:
     if fragment:
      style=next((z for z in native.get('paragraphStyles',[]) if fragment in z['text']),{});nodes.append(dict(kind='text',text=fragment,x=offset,y=li*fs*1.22,fontSize=fs,fontWeight=style.get('fontWeight',native.get('fontWeight',400)),fill='#111111',afterFill=style.get('fill',native.get('fill','#103758')),id=f'a{ci}-t{len(nodes)}',sourceRef=dict(component=item['id'],path='text')))
      offset+=fonts.load_font('黑体',round(fs)).getlength(fragment)
  elif value['kind']=='table':
   import semantic_scene as semantic
   nodes=semantic.compile_value(value,Image.new('RGB',(1,1)),(w,h),Path('.'),f'auto{ci}',item['id'])
  elif value['kind']=='image':
   if not native.get('imageData'):raise ValueError('自动照片缺少独立素材；请重新读取课件')
   nodes=[dict(kind='image',x=0,y=0,w=w,h=h,imageData=native['imageData'],fit='contain')]
  else:raise ValueError('自动布局当前支持文字、完整表格与独立照片，请重新读取生成方案')
  if native.get('background'):nodes.insert(0,dict(kind='rect',x=0,y=0,w=w,h=h,fill=native['background']))
  for ni,n in enumerate(nodes):n.setdefault('id',f'a{ci}-n{ni}')
  groups.append(dict(id=item['id'],flowParent=native.get('flowGroup',item['id']),size=[w,h],before=[x,y,1],after=[x,y,1],index=min(ci,39),nodes=nodes))
 # Decorative source rectangles remain target-only; never carry them into the ordinary page.
 for d in sorted(items[0].get('autoNative',{}).get('decor',[]),key=lambda z:z['box'][2]*z['box'][3]) if items else []:
  x,y,w,h=d['box']
  if d.get('kind')=='line':
   groups[0]['nodes'].insert(0,dict(kind='line',x=x*W-groups[0]['after'][0],y=y*H-groups[0]['after'][1],points=[[a*W,b*H] for a,b in d['points']],fill=d['fill'],stroke=d['fill'],lineWidth=1.5 if not d.get('closed',True) else .5,closed=d.get('closed',True),before=dict(opacity=0)))
  else:groups[0]['nodes'].insert(0,dict(kind='rect',x=x*W-groups[0]['after'][0],y=y*H-groups[0]['after'][1],w=w*W,h=h*H,fill=d['fill'],before=dict(opacity=0)))
 # Keep decorations in a page-sized group so their geometry does not alter text containers.
 decorNodes=[n for n in groups[0]['nodes'] if n.get('before',{}).get('opacity')==0] if groups else []
 if decorNodes:
  ox,oy,_=groups[0]['after'];groups[0]['nodes']=[n for n in groups[0]['nodes'] if n not in decorNodes]
  for n in decorNodes:n['x']+=ox;n['y']+=oy
  groups.append(dict(id='auto-decoration',size=[W,H],before=[0,0,1],after=[0,0,1],index=0,nodes=decorNodes))
 # Independent plain reading layout: continuous text columns, tables as basic grids, photos below.
 texts=[(g,n) for g in groups for n in g['nodes'] if n['kind']=='text'];photos=[(g,n) for g in groups for n in g['nodes'] if n['kind']=='image'];tableGroups=[g for g,item in zip(groups,items) if item['plain']['kind']=='table']
 titleIds={item['id'] for item in items if item['plain'].get('role')=='title'}
 title=next(((g,n) for g,n in texts if g['id'] in titleIds),None)
 tableIds={g['id'] for g in tableGroups}
 body=[(g,n) for g,n in texts if g['id'] not in tableIds and g['id'] not in titleIds]
 photoH=190 if photos else 0;bottom=H-35-photoH;cols=1 if tableGroups else 3 if sum(len(n['text']) for g,n in body)>250 else 2;startX=W*.52 if tableGroups else 40;cw=(W-startX-60)/cols
 for ti,g in enumerate(tableGroups):
  slotW=W*.46 if body else W-80;slotH=(bottom-150)/max(1,len(tableGroups));factor=min(slotW/g['size'][0],slotH/g['size'][1]);tx=40;ty=130+ti*slotH
  for n in g['nodes']:
   b=dict(x=tx+n['x']*factor,y=ty+n['y']*factor,opacity=1)
   if n['kind']=='text':b.update(fontSize=n['fontSize']*factor,fontWeight=400,fill='#111111')
   elif n['kind']=='rect':b.update(w=n['w']*factor,h=n['h']*factor,fill='#ffffff',stroke='#777777',lineWidth=.8,radius=0)
   n['before']=b
 for fs in [26,24,22,20,18,16,14]:
  col=0;xx=startX;yy=130;placements=[];lastGroup=None;fits=True
  for g,n in body:
   key=g.get('flowParent',g['id'])
   if key!=lastGroup:
    if lastGroup:yy+=fs*1.15;xx=startX+col*cw
    lastGroup=key
   width=fonts.load_font('黑体',fs).getlength(n['text'])
   if width>cw-25:fits=False;break
   if xx+width>startX+(col+1)*cw-20:xx=startX+col*cw;yy+=fs*1.3
   if yy+fs*1.25>bottom:
    col+=1;xx=startX+col*cw;yy=130
   if col>=cols:fits=False;break
   placements.append((n,dict(x=xx,y=yy,fontSize=fs,fontWeight=400,fill='#111111',opacity=1)));xx+=width+fs*.6
  if fits:break
 else:raise ValueError('普通版完整内容在可读字号下放不下，请分成多页后重试')
 for n,b in placements:n['before']=b
 if title:
  titleNodes=[n for g,n in texts if g['id'] in titleIds];xx=40;yy=25
  for n in titleNodes:
   width=fonts.load_font('黑体',34).getlength(n['text'])
   if xx+width>W-40:xx=40;yy+=42
   n['before']=dict(x=xx,y=yy,fontSize=34,fontWeight=400,fill='#111111');xx+=width+12
 for i,(g,n) in enumerate(photos):
  columns=min(6,len(photos));rows=math.ceil(len(photos)/columns);slotW=(W-80)/columns;slotH=photoH/max(1,rows)
  n['before']=dict(x=40+i%columns*slotW,y=H-photoH-25+i//columns*slotH,w=slotW-12,h=slotH-12)
 for g in groups:
  for n in g['nodes']:
   if n['kind'] in ['rect','line'] and 'before' not in n:n['before']=dict(opacity=0)
 return dict(width=W,height=H,groups=[g for g in groups if g['id']=='auto-decoration']+[g for g in groups if g['id']!='auto-decoration'])

def pack_items(items):
 """Bound renderer group count without dropping any shape; nearby objects share a semantic container."""
 items=photo_semantics(items)
 if len(items)<=39:return items
 flattened=[]
 for item in items:
  if item['plain']['kind']!='group':flattened.append(item);continue
  for i,(part,meta) in enumerate(zip(item['plain']['parts'],item['autoNative']['members'])):
   flattened.append(dict(id=item['id']+'-leaf-'+str(i),name=item['name'],plain={k:v for k,v in part.items() if k!='box'},target=meta['box'],before=item['before'],autoNative=meta))
 items=flattened
 buckets={};titles=[]
 for item in items:
  if item['plain'].get('role')=='title' and len(titles)<3:titles.append(item);continue
  x,y,w,h=item['target'];key=(min(3,int((x+w/2)*4)),min(2,int((y+h/2)*3)));buckets.setdefault(key,[]).append(item)
 result=list(titles)
 for key,values in sorted(buckets.items(),key=lambda kv:(kv[0][1],kv[0][0])):
  values.sort(key=lambda z:(round(z['target'][1],2),z['target'][0]))
  for start in range(0,len(values),20):
   batch=values[start:start+20];x=min(z['target'][0] for z in batch);y=min(z['target'][1] for z in batch);w=max(z['target'][0]+z['target'][2] for z in batch)-x;h=max(z['target'][1]+z['target'][3] for z in batch)-y
   parts=[]
   for z in batch:
    a,b,c,d=z['target'];parts.append(dict(z['plain'],box=[(a-x)/w,(b-y)/h,c/w,d/h]))
   result.append(dict(id='auto-group-'+str(len(result)),name='自动归组 '+str(len(result)+1),text='',target=[x,y,w,h],before=[.04,.04,.92,.92],plain=dict(kind='group',parts=parts),autoNative=dict(members=[z['autoNative'] for z in batch],canvas=items[0]['autoNative'].get('canvas',[1600,900]))))
 decor=items[0]['autoNative'].get('decor')
 if decor and result:result[0]['autoNative']['decor']=decor
 return result

def from_raster(source,recognized=None):
 """Local OCR plus conservative textured-photo regions; never use a whole-page image node."""
 import tempfile
 import numpy as np
 import cv2
 try:import local_ocr
 except ImportError:return [],[dict(code='ocr-unavailable',message='本机尚未配置本地OCR；请安装受支持的本地识别环境后重新读取，不能把整页截图作普通版')]
 if recognized is None:
  with tempfile.TemporaryDirectory() as d:
   path=Path(d)/'page.png';source.save(path);result=recognize_page(path)
 else:result=recognized
 if not result.get('available'):return [],[dict(code='ocr-unavailable',message='本地OCR不可用：'+str(z)) for z in result.get('issues',[])] or [dict(code='ocr-unavailable',message='本地OCR不可用，请检查识别环境')]
 lines=[z for z in result.get('lines',[]) if z.get('text','').strip()]
 if not lines:return [],[dict(code='ocr-no-text',message='本页未读出文字，请使用清晰原页或导入可编辑课件')]
 if len(lines)>400:return [],[dict(code='ocr-dense',message='本页超过400行识别文字，未截断；请拆页后重试')]
 W,H=source.size;scale=1600/W;array=np.array(source.convert('RGB'));gray=cv2.cvtColor(array,cv2.COLOR_RGB2GRAY);hsv=cv2.cvtColor(array,cv2.COLOR_RGB2HSV);textmask=np.zeros((H,W),np.uint8)
 for z in lines:
  x,y,w,h=z['box'];l=max(0,round(x*W)-2);t=max(0,round(y*H)-2);r=min(W,round((x+w)*W)+2);b=min(H,round((y+h)*H)+2);textmask[t:b,l:r]=255
 # Detect natural texture after removing OCR glyph regions. Flat chart bars and panels are not photos.
 g=gray.astype(np.float32);variance=cv2.boxFilter(g*g,-1,(9,9))-cv2.boxFilter(g,-1,(9,9))**2
 rich=((variance>150)&(hsv[:,:,1]>14)).astype(np.float32)
 fraction=cv2.boxFilter(rich,-1,(15,15))
 texture=((fraction>.48)&((textmask==0)|(fraction>.82))).astype(np.uint8)*255
 texture=cv2.morphologyEx(texture,cv2.MORPH_CLOSE,cv2.getStructuringElement(cv2.MORPH_RECT,(5,5)))
 candidates=[]
 for contour in cv2.findContours(texture,cv2.RETR_EXTERNAL,cv2.CHAIN_APPROX_SIMPLE)[0]:
  x,y,w,h=cv2.boundingRect(contour);area=w*h
  if not .0015*W*H<area<.28*W*H or min(w,h)<min(W,H)*.035 or not .22<w/h<4.8:continue
  crop=gray[y:y+h,x:x+w];edge=cv2.Canny(crop,60,140);density=(edge>0).mean();textCoverage=(textmask[y:y+h,x:x+w]>0).mean();sat=hsv[y:y+h,x:x+w,1]
  if density<.055 or textCoverage>.35 or np.std(crop)<28 or (sat>35).mean()<.12:continue
  candidates.append([x,y,w,h])
 candidates=photo_regions(array,textmask,lines,result.get('rectangles',[]))
 # Merge overlapping fragments of the same photo, but leave separated photos independent.
 changed=True
 while changed:
  changed=False
  for i in range(len(candidates)):
   for j in range(i):
    a,b=candidates[i],candidates[j];ix=max(0,min(a[0]+a[2],b[0]+b[2])-max(a[0],b[0]));iy=max(0,min(a[1]+a[3],b[1]+b[3])-max(a[1],b[1]))
    if ix*iy>min(a[2]*a[3],b[2]*b[3])*.3:
     x=min(a[0],b[0]);y=min(a[1],b[1]);r=max(a[0]+a[2],b[0]+b[2]);bt=max(a[1]+a[3],b[1]+b[3]);candidates[j]=[x,y,r-x,bt-y];candidates.pop(i);changed=True;break
   if changed:break
 embeddedPhotoText=[];kept=[]
 for z in lines:
  x,y,w,h=z['box'];inside=any(a<=x*W and b<=y*H and (x+w)*W<=a+c+3 and (y+h)*H<=b+d+3 for a,b,c,d in candidates)
  if inside:embeddedPhotoText.append(z['text'])
  else:kept.append(z)
 lines=kept
 # Recover ruled tables as actual cell data, preserving row/column relations.
 binary=cv2.threshold(gray,190,255,cv2.THRESH_BINARY_INV)[1]
 horizontal=cv2.morphologyEx(binary,cv2.MORPH_OPEN,cv2.getStructuringElement(cv2.MORPH_RECT,(max(24,W//22),1)))
 vertical=cv2.morphologyEx(binary,cv2.MORPH_OPEN,cv2.getStructuringElement(cv2.MORPH_RECT,(1,max(18,H//25))))
 grid=cv2.add(horizontal,vertical);tables=[];consumed=set()
 def bands(values):
  result=[]
  for v in values:
   if not result or v-result[-1][-1]>3:result.append([int(v)])
   else:result[-1].append(int(v))
  return [sum(z)/len(z) for z in result]
 for contour in cv2.findContours(grid,cv2.RETR_EXTERNAL,cv2.CHAIN_APPROX_SIMPLE)[0]:
  x,y,w,h=cv2.boundingRect(contour)
  if w<W*.15 or h<H*.12:continue
  xs=bands(np.where((vertical[y:y+h,x:x+w]>0).mean(axis=0)>.50)[0]);ys=bands(np.where((horizontal[y:y+h,x:x+w]>0).mean(axis=1)>.50)[0])
  if not 3<=len(xs)<=13 or not 3<=len(ys)<=31:continue
  rows=[['' for _ in xs[:-1]] for _ in ys[:-1]];members=[]
  for li,z in enumerate(lines):
   zx,zy,zw,zh=z['box'];cx=(zx+zw/2)*W-x;cy=(zy+zh/2)*H-y
   if not xs[0]<=cx<=xs[-1] or not ys[0]<=cy<=ys[-1]:continue
   col=min(len(xs)-2,max(0,int(np.searchsorted(xs,cx)-1)));row=min(len(ys)-2,max(0,int(np.searchsorted(ys,cy)-1)));rows[row][col]+=z['text'];members.append(li)
  if sum(bool(c) for row in rows for c in row)<4:continue
  tables.append(dict(id='table-'+str(len(tables)),name='本地表格 '+str(len(tables)+1),text='',target=[x/W,y/H,w/W,h/H],before=[.04,.04,.92,.92],plain=dict(kind='table',rows=rows,columns=[xs[i+1]-xs[i] for i in range(len(xs)-1)]),autoNative=dict(kind='table',box=[x/W,y/H,w/W,h/H],canvas=[1600,H*scale],detection='ruled-grid',sourceLineIds=members)))
  consumed.update(members)
 items=list(tables)
 lines=[z for i,z in enumerate(lines) if i not in consumed]
 lines.sort(key=lambda z:(round(z['box'][1]*H/18),z['box'][0]))
 for i,z in enumerate(lines):
  box=z['box'];x,y,w,h=box;txt=z['text'];left,top,right,bt=round(x*W),round(y*H),round((x+w)*W),round((y+h)*H);pixel=source.crop((left,top,right,bt));pix=np.array(pixel)
  ring=array[max(0,top-4):min(H,bt+4),max(0,left-4):min(W,right+4)].copy();border=np.concatenate([ring[:3].reshape(-1,3),ring[-3:].reshape(-1,3),ring[:,:3].reshape(-1,3),ring[:,-3:].reshape(-1,3)]) if ring.size else np.array([[255,255,255]])
  background=np.median(border,axis=0);distance=np.linalg.norm(pix.astype(float)-background,axis=2);foregroundPixels=pix[distance>=max(35,float(np.percentile(distance,75)))];foreground=np.median(foregroundPixels,axis=0) if len(foregroundPixels) else [30,30,30];colourHex='#'+''.join(f'{int(v):02x}' for v in foreground);fs=max(14,h*H*scale*1.05)

  role='title' if y<.10 and fs>26 and len(txt)<100 else 'body'
  items.append(dict(id='ocr-'+str(i),name='识别文字 '+str(i+1),text='',target=[max(0,x),max(0,y),min(1-x,w+.008),min(1-y,h*1.3)],before=[.04,.04,.92,.92],plain=dict(kind='text',text=txt,role=role),autoNative=dict(kind='text',box=box,fontSize=fs,fontWeight=700 if role=='title' else 400,fill=colourHex,canvas=[1600,H*scale],ocr=dict(confidence=z.get('confidence'),backend=result.get('backend'),alternatives=z.get('alternatives',[])))))
 for i,(x,y,w,h) in enumerate(sorted(candidates,key=lambda b:(b[1],b[0]))):
  # Only textured photo pixels, never a page-sized or text/table region.
  bitmap=source.crop((x,y,x+w,y+h));buf=io.BytesIO();bitmap.save(buf,format='PNG');box=[x/W,y/H,w/W,h/H]
  items.append(dict(id='photo-'+str(i),name='本地照片候选 '+str(i+1),text='',target=box,before=[.04,.04,.92,.92],plain=dict(kind='image',assetRole='photo'),autoNative=dict(kind='image',box=box,imageData=base64.b64encode(buf.getvalue()).decode(),canvas=[1600,H*scale],detection='texture-region; preview must confirm photo bounds')))
 if items:
  vectors,vectorIssues=recover_vectors(array,textmask,candidates,items)
  if vectorIssues:return items,vectorIssues
  items[0]['autoNative']['decor']=vectors
  items[0]['autoNative']['rasterAudit']=dict(backend=result.get('backend'),ocrLines=len(lines),photoCandidates=len(candidates),embeddedPhotoText=embeddedPhotoText,lowConfidenceLines=sum(z.get('confidence',1)<.6 for z in lines),sourceSize=[W,H],comparison=result.get('comparison'),tableCount=len(tables),vectorShapes=len(vectors),boundary='OCR and photo candidates require visual confirmation; no source-slide completeness guarantee')
 return items,[]


def recover_vectors(array,textmask,photos,text_items=()):
 """Detect clean rectangles, connectors and straight structural lines; never trace residual glyph pixels."""
 import numpy as np,cv2
 H,W=array.shape[:2];hsv=cv2.cvtColor(array,cv2.COLOR_RGB2HSV);valid=np.ones((H,W),dtype=bool)
 for x,y,w,h in photos:valid[max(0,y-3):min(H,y+h+3),max(0,x-3):min(W,x+w+3)]=False
 shapes=[];rectangles=[]
 for bucket in range(15):
  selected=(hsv[:,:,0]//12==bucket)&(hsv[:,:,1]>65)&(hsv[:,:,2]>55)&valid
  if selected.sum()<W*H*.0003:continue
  mask=selected.astype(np.uint8)*255;mask=cv2.morphologyEx(mask,cv2.MORPH_CLOSE,np.ones((5,5),np.uint8))
  for c in cv2.findContours(mask,cv2.RETR_EXTERNAL,cv2.CHAIN_APPROX_SIMPLE)[0]:
   area=cv2.contourArea(c);x,y,w,h=cv2.boundingRect(c)
   if area<W*H*.0003:continue
   inkPixels=array[y:y+h,x:x+w][selected[y:y+h,x:x+w]]
   if not len(inkPixels):continue
   rgb=np.median(inkPixels,axis=0).astype(int);fill='#'+''.join(f'{v:02x}' for v in rgb)
   foreground=False
   for item in text_items:
    if item['plain']['kind']!='text':continue
    a,b,cw,ch=item['target'];a*=W;b*=H;cw*=W;ch*=H;overlap=max(0,min(x+w,a+cw)-max(x,a))*max(0,min(y+h,b+ch)-max(y,b));value=item['autoNative'].get('fill','#000000');fg=np.array([int(value[i:i+2],16) for i in (1,3,5)])
    if overlap/min(w*h,max(1,cw*ch))>.45 and np.linalg.norm(rgb-fg)<100:foreground=True;break
   if foreground:continue
   if area/(w*h)>.70 and (min(w,h)>20 or max(w,h)>W*.2):
    # Clean source panel/title bar. Glyph cutouts and gradients never become polygon holes.
    box=[x/W,y/H,w/W,h/H];shapes.append(dict(box=box,fill=fill));rectangles.append(box);continue
   if .35<w/max(1,h)<3 and max(w,h)<W*.14 and area/(w*h)>.30 and (textmask[y:y+h,x:x+w]>0).mean()<.12:
    left=(mask[y:y+h,x:x+max(1,w//5)]>0).sum();right=(mask[y:y+h,x+w-max(1,w//5):x+w]>0).sum()
    if abs(left-right)>area*.08 and w>h*.6:
     # Canonical arrow, inferred direction, keeping the detected source bounds.
     points=[[0,h*.28],[w*.55,h*.28],[w*.55,0],[w,h*.5],[w*.55,h],[w*.55,h*.72],[0,h*.72]]
     if left<right:points=[[w-a,b] for a,b in points]
     shapes.append(dict(kind='line',box=[x/W,y/H,w/W,h/H],points=[[a/W,b/H] for a,b in points],fill=fill))
 # Black/grey diagrams and table/axis relationships are line primitives, not a flattened crop.
 gray=cv2.cvtColor(array,cv2.COLOR_RGB2GRAY);lineMask=((gray<170)&(hsv[:,:,1]<60)&valid&(textmask==0)).astype(np.uint8)*255
 lineMask=cv2.morphologyEx(lineMask,cv2.MORPH_OPEN,np.ones((1,3),np.uint8))
 lines=cv2.HoughLinesP(lineMask,1,np.pi/180,threshold=45,minLineLength=max(45,W//30),maxLineGap=5)
 if lines is not None:
  for x1,y1,x2,y2 in lines.reshape(-1,4)[:120]:
   if min(x1,x2)==max(x1,x2) and abs(y1-y2)<H*.08:continue
   x=min(x1,x2);y=min(y1,y2);w=abs(x2-x1);h=abs(y2-y1)
   shapes.append(dict(kind='line',box=[x/W,y/H,max(1,w)/W,max(1,h)/H],points=[[(x1-x)/W,(y1-y)/H],[(x2-x)/W,(y2-y)/H]],fill='#58636e',closed=False))
 return shapes,[]


def attach_structures(scene):
 """Bind detected clean plates to semantic content; only global connectors stay on the page."""
 decor=next((g for g in scene['groups'] if g['id']=='auto-decoration'),None)
 if not decor:return scene
 content=[g for g in scene['groups'] if g is not decor];assigned={g['id']:[] for g in content};remaining=[]
 for n in decor['nodes']:
  x,y=n['x'],n['y'];w=n.get('w',max((p[0] for p in n.get('points',[[0,0]])),default=0));h=n.get('h',max((p[1] for p in n.get('points',[[0,0]])),default=0));options=[]
  for g in content:
   gx,gy,_=g['after'];gw,gh=g['size'];intersection=max(0,min(x+w,gx+gw)-max(x,gx))*max(0,min(y+h,gy+gh)-max(y,gy))
   if w*h and intersection/min(w*h,gw*gh)>.2:options.append((intersection/max(1,w*h+gw*gh-intersection),g))
  if options and n['kind']=='rect':assigned[max(options,key=lambda z:z[0])[1]['id']].append(n)
  else:remaining.append(n)
 for g in content:
  plates=assigned[g['id']]
  if not plates:continue
  ox,oy,_=g['after'];x=min([ox]+[n['x'] for n in plates]);y=min([oy]+[n['y'] for n in plates]);right=max([ox+g['size'][0]]+[n['x']+n['w'] for n in plates]);bt=max([oy+g['size'][1]]+[n['y']+n['h'] for n in plates])
  for n in g['nodes']:n['x']+=ox-x;n['y']+=oy-y
  for n in plates:n['x']-=x;n['y']-=y
  g['nodes']=plates+g['nodes'];g['after']=g['before']=[x,y,1];g['size']=[right-x,bt-y]
 decor['nodes']=remaining
 if not remaining:scene['groups'].remove(decor)
 return scene

def photo_regions(array,textmask,lines,rectangles):
 import cv2,numpy as np
 H,W=array.shape[:2];gray=cv2.cvtColor(array,cv2.COLOR_RGB2GRAY);hsv=cv2.cvtColor(array,cv2.COLOR_RGB2HSV);g=gray.astype(np.float32);var=cv2.boxFilter(g*g,-1,(9,9))-cv2.boxFilter(g,-1,(9,9))**2;rich=((var>150)&(hsv[:,:,1]>14)).astype(np.float32);fraction=cv2.boxFilter(rich,-1,(15,15));mask=((fraction>.48)&(textmask==0)).astype(np.uint8)
 proposals=[]
 # Repeated photo rows are separated by texture projections, without borrowing card titles/body.
 rows=(mask.mean(axis=1)>.055).astype(np.uint8).reshape(-1,1);rows=cv2.morphologyEx(rows,cv2.MORPH_CLOSE,np.ones((11,1),np.uint8))
 for c in cv2.findContours(rows,cv2.RETR_EXTERNAL,cv2.CHAIN_APPROX_SIMPLE)[0]:
  _,y,_,h=cv2.boundingRect(c)
  if h<max(45,H*.05):continue
  cols=(mask[y:y+h].mean(axis=0)>.25).astype(np.uint8).reshape(1,-1);cols=cv2.morphologyEx(cols,cv2.MORPH_CLOSE,np.ones((1,5),np.uint8))
  for cc in cv2.findContours(cols,cv2.RETR_EXTERNAL,cv2.CHAIN_APPROX_SIMPLE)[0]:
   x,_,w,_=cv2.boundingRect(cc)
   if w>max(50,W*.025):proposals.append([x,y,w,h,'texture-band'])
 edges=cv2.morphologyEx(cv2.Canny(gray,40,120),cv2.MORPH_CLOSE,np.ones((5,5),np.uint8))
 for c in cv2.findContours(edges,cv2.RETR_LIST,cv2.CHAIN_APPROX_SIMPLE)[0]:
  x,y,w,h=cv2.boundingRect(c);poly=cv2.approxPolyDP(c,.025*cv2.arcLength(c,True),True)
  if 4<=len(poly)<=8 and min(w,h)>max(45,min(W,H)*.05) and cv2.contourArea(c)/(w*h)>.8:proposals.append([x,y,w,h,'edge-rectangle'])
 for z in rectangles:
  x,y,w,h=z['box'];proposals.append([round(x*W),round(y*H),round(w*W),round(h*H),'vision-rectangle'])
 accepted=[]
 for x,y,w,h,kind in proposals:
  x=max(0,x);y=max(0,y);w=min(W-x,w);h=min(H-y,h)
  if not .002*W*H<w*h<.20*W*H or not .35<w/max(1,h)<2.6 or min(w,h)<45:continue
  crop=gray[y:y+h,x:x+w];coverage=(textmask[y:y+h,x:x+w]>0).mean();texture=(fraction[y:y+h,x:x+w]>.48).mean();edgeDensity=(cv2.Canny(crop,60,140)>0).mean();white=(crop>235).mean()
  inside=[z for z in lines if x<=(z['box'][0]+z['box'][2]/2)*W<=x+w and y<=(z['box'][1]+z['box'][3]/2)*H<=y+h and z['box'][3]*H>14]
  if len(inside)>=3 or coverage>.3 or texture<.2 or edgeDensity<.05 or np.std(crop)<28 or white>.62:continue
  # Prefer explicit edge rectangles over loose texture envelopes.
  score=(2 if kind=='edge-rectangle' else 1 if kind=='vision-rectangle' else 0)+texture-coverage
  accepted.append((score,[x,y,w,h]))
 result=[]
 for score,b in sorted(accepted,reverse=True):
  x,y,w,h=b
  if any(max(0,min(x+w,a+c)-max(x,a))*max(0,min(y+h,q+d)-max(y,q))>min(w*h,c*d)*.45 for a,q,c,d in result):continue
  result.append(b)
 # Repeated numbered workflow headings expose semantic columns. When a textured photo row
 # is present, recover all aligned slots rather than treating a card/header as one bitmap.
 anchors=[]
 for line in lines:
  match=re.match(r'^\s*(\d{1,2})(?=\s|[\u3400-\u9fff]|$)',line['text'])
  if match:
   x,y,w,h=line['box'];anchors.append(dict(number=int(match.group(1)),x=x*W,y=y*H,h=h*H))
 rowsByY=[]
 for anchor in sorted(anchors,key=lambda z:z['y']):
  if rowsByY and abs(anchor['y']-sum(z['y'] for z in rowsByY[-1])/len(rowsByY[-1]))<18:rowsByY[-1].append(anchor)
  else:rowsByY.append([anchor])
 for row in rowsByY:
  row.sort(key=lambda z:z['x'])
  if len(row)<3 or any(row[i+1]['number']-row[i]['number']!=1 for i in range(len(row)-1)):continue
  spacing=float(np.median([row[i+1]['x']-row[i]['x'] for i in range(len(row)-1)]));hh=float(np.median([z['h'] for z in row]));baseY=max(z['y']+z['h'] for z in row)
  below=[b for b in result if baseY-10<b[1]<baseY+hh*3 and b[2]<spacing*1.1 and b[3]>hh*2]
  if len(below)<2:continue
  yy=float(np.median([b[1] for b in below]));height=float(np.median([b[3] for b in below]));width=min(spacing-hh*1.2,float(np.median([b[2] for b in below if b[2]>spacing*.65] or [spacing-hh*1.2])))
  if yy+height>H or width<50:continue
  result=[b for b in result if not (b[1]<yy+height and b[1]+b[3]>yy)]
  for anchor in row:result.append([max(0,round(anchor['x']-hh*.12)),round(yy),round(width),round(height)])
 return result

def expand_text_images(items,cache):
 """Inspect large embedded pictures locally; screenshots cannot silently pass as photos."""
 import hashlib,json,tempfile
 import local_ocr
 cache=Path(cache);cache.mkdir(parents=True,exist_ok=True);output=[];issues=[]
 for item in items:
  meta=item.get('autoNative',{});raw=meta.get('imageData')
  if item['plain']['kind']!='image' or not raw:output.append(item);continue
  im=Image.open(io.BytesIO(base64.b64decode(raw))).convert('RGB')
  if im.width<240 or im.height<100:meta['imageInspection']='small-asset';output.append(item);continue
  image_key=hashlib.sha256(base64.b64decode(raw)).hexdigest();runtime_key=b''.join(Path(__file__).with_name(n).read_bytes() for n in ['auto_components.py','local_ocr.py','local_ocr.swift','semantic_scene.py','plain_components.py','component_video.py'] if Path(__file__).with_name(n).exists());key=hashlib.sha256(image_key.encode()+runtime_key).hexdigest();record=cache/(key+'.json')
  if record.exists():ocr=json.loads(record.read_text())
  else:
   with tempfile.TemporaryDirectory() as d:
    p=Path(d)/'image.png';im.save(p);ocr=recognize_page(p)
   record.write_text(json.dumps(ocr,ensure_ascii=False))
  if not ocr.get('available'):issues.append(dict(code='embedded-ocr-unavailable',message='大图片内文字检查失败，请恢复本地OCR后重新读取',location=item.get('name')));output.append(item);continue
  lines=ocr.get('lines',[]);chars=sum(len(z.get('text','')) for z in lines)
  if len(lines)<3 or chars<25:meta['imageInspection']=dict(kind='photo-or-logo',ocrLines=len(lines));output.append(item);continue
  children,errors=from_raster(im,ocr)
  if errors:issues.extend(dict(z,location=item.get('name')) for z in errors);output.append(item);continue
  x,y,w,h=item['target'];canvas=meta.get('canvas',[1600,900])
  for j,child in enumerate(children):
   cx,cy,cw,ch=child['target'];child['target']=[x+cx*w,y+cy*h,cw*w,ch*h];child['id']=item['id']+'-embedded-'+str(j);child['name']=item.get('name','图片')+' · '+child['name'];cm=child['autoNative'];cm['box']=child['target'];cm['canvas']=canvas
   if 'fontSize' in cm:cm['fontSize']*=w
   for d in cm.get('decor',[]):
    a,b,c,e=d['box'];d['box']=[x+a*w,y+b*h,c*w,e*h]
    if 'points' in d:d['points']=[[px*w,py*h] for px,py in d['points']]
   cm['embeddedOrigin']=dict(id=item['id'],sha256=image_key,ocrLines=len(lines));output.append(child)
 return output,issues

def combine_members(batch,name,ident):
 x=min(z['target'][0] for z in batch);y=min(z['target'][1] for z in batch);w=max(z['target'][0]+z['target'][2] for z in batch)-x;h=max(z['target'][1]+z['target'][3] for z in batch)-y
 parts=[]
 for z in batch:
  a,b,c,d=z['target'];parts.append(dict(z['plain'],box=[(a-x)/w,(b-y)/h,c/w,d/h]))
 return dict(id=ident,name=name,text='',target=[x,y,w,h],before=[.04,.04,.92,.92],plain=dict(kind='group',parts=parts),autoNative=dict(members=[z['autoNative'] for z in batch],canvas=batch[0]['autoNative'].get('canvas',[1600,900])))

def photo_semantics(items):
 photos=[z for z in items if z['plain']['kind']=='image'];buckets={p['id']:[p] for p in photos};rest=[]
 if len(photos)<2:return items
 for item in items:
  if item['plain']['kind']=='image':continue
  if item['plain']['kind']!='text' or item['plain'].get('role')=='title':rest.append(item);continue
  x,y,w,h=item['target'];cx=x+w/2;cy=y+h/2;candidates=[]
  eligible=[p for p in photos if p['target'][1]-p['target'][3]*.6<=cy]
  rowStart=max((p['target'][1]-p['target'][3]*.6 for p in eligible),default=-1)
  for photo in photos:
   a,b,c,d=photo['target'];px=a+c/2;py=b+d/2
   if rowStart>=0 and abs((b-d*.6)-rowStart)>.05:continue
   if w>c*2.6 or x<a-.045 or abs(cx-px)>c*1.7 or abs(cy-py)>d*1.6:continue
   candidates.append((abs(cx-px)+abs(cy-py)*.45,photo))
  if not candidates:rest.append(item);continue
  photo=min(candidates,key=lambda z:z[0])[1];buckets[photo['id']].append(item)
 for i,batch in enumerate(buckets.values()):
  batch.sort(key=lambda z:(round(z['target'][1],2),z['target'][0]))
  for j in range(0,len(batch),20):rest.append(combine_members(batch[j:j+20],'自动图文语义组 '+str(i+1),'semantic-'+str(i)+'-'+str(j)))
 rest.sort(key=lambda z:(round(z['target'][1]/.08),z['target'][0]))
 decor=items[0].get('autoNative',{}).get('decor')
 if decor and rest:rest[0]['autoNative']['decor']=decor
 return rest

def recognize_page(path):
 """Use installed local text models; Vision supplies independent geometry/evidence on macOS."""
 import local_ocr,importlib.util,sys
 vision=local_ocr.recognize(path) if sys.platform=='darwin' else dict(available=False,lines=[],rectangles=[],issues=[])
 rapid=local_ocr.recognize(path,backend='rapidocr') if importlib.util.find_spec('rapidocr') else dict(available=False,lines=[],issues=[])
 chosen=rapid if rapid.get('available') and rapid.get('lines') else vision
 if not chosen.get('available'):return dict(available=False,lines=[],issues=rapid.get('issues',[])+vision.get('issues',[]) or ['No local OCR backend is configured'])
 chosen=dict(chosen,rectangles=vision.get('rectangles',[]));disagreements=[]
 def norm(text):return ''.join(re.findall(r'[\u3400-\u9fffA-Za-z0-9]',text))
 if chosen is not vision and vision.get('available'):
  for z in chosen.get('lines',[]):
   x,y,w,h=z['box'];alternatives=[];matches=[]
   for other in vision.get('lines',[]):
    a,b,c,d=other['box'];overlap=max(0,min(x+w,a+c)-max(x,a))*max(0,min(y+h,b+d)-max(y,b))
    if overlap>min(w*h,c*d)*.4 and abs(y+h/2-b-d/2)<max(h,d)*.6:alternatives.append(other['text']);matches.append(other)
   if matches and norm(z['text'])==norm(''.join(alternatives)):
    boxes=[m['box'] for m in matches];left=min(b[0] for b in boxes);top=min(b[1] for b in boxes);right=max(b[0]+b[2] for b in boxes);bottom=max(b[1]+b[3] for b in boxes);z['rawBox']=z['box'];z['box']=[left,top,right-left,bottom-top];z['geometryBackend']='apple_vision'
   if alternatives and norm(z['text'])!=norm(''.join(alternatives)):
    z['alternatives']=alternatives;disagreements.append(dict(text=z['text'],alternatives=alternatives,box=z['box']))
 chosen['comparison']=dict(otherBackend=vision.get('backend'),disagreements=disagreements,requiresVisualReview=True)
 return chosen

def analysis_fingerprint():
 import hashlib,json
 from importlib.metadata import version,PackageNotFoundError
 root=Path(__file__).parent;values={n:hashlib.sha256((root/n).read_bytes()).hexdigest() for n in ['auto_components.py','local_ocr.py','local_ocr.swift','semantic_scene.py','plain_components.py','component_video.py'] if (root/n).exists()}
 for name in ['rapidocr','onnxruntime','opencv-python-headless']:
  try:values[name]=version(name)
  except PackageNotFoundError:values[name]='unavailable'
 return hashlib.sha256(json.dumps(values,sort_keys=True).encode()).hexdigest()
