"""Create fictional, redistributable native + screenshot PPT example."""
from pathlib import Path
import sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'engine'))
from xhs_engine import fonts
from PIL import Image,ImageDraw
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
out=ROOT/'examples'/'agent-video';out.mkdir(parents=True,exist_ok=True)
prs=Presentation();prs.slide_width=Inches(12);prs.slide_height=Inches(6.75)
slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string('F0F4F8')
items=[('title','项目复盘 · 从想法到交付',[.05,.06,.9,.15],'132C4A','FFFFFF',28),('step1','01  明确目标\n整理需求与验收标准',[.06,.30,.27,.52],'D8E8F2','132C4A',22),('step2','02  制作初稿\n完成结构与内容检查',[.365,.30,.27,.52],'D1EBE7','132C4A',22),('step3','03  核对交付\n预览修正并保存成品',[.67,.30,.27,.52],'F4E8D3','132C4A',22)]
for name,text,r,bg,fg,size in items:
 x,y,w,h=r;s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x*12),Inches(y*6.75),Inches(w*12),Inches(h*6.75));s.name=name;s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(bg);s.line.fill.background();s.text=text
 for p in s.text_frame.paragraphs:p.font.name='PingFang SC';p.font.size=Pt(size);p.font.color.rgb=RGBColor.from_string(fg)
im=Image.new('RGB',(1600,900),'#f0f4f8');d=ImageDraw.Draw(im)
regions=[('title','交付检查 · 三个关键动作',[.05,.06,.9,.15],'#132c4a'),('check1','核对内容\n标题、数字、单位一致',[.06,.30,.27,.52],'#d8e8f2'),('check2','查看画面\n组件完整、文字清楚',[.365,.30,.27,.52],'#d1ebe7'),('check3','保存结果\n视频、方案、来源齐全',[.67,.30,.27,.52],'#f4e8d3')]
for name,text,r,bg in regions:
 x,y,w,h=r;box=(int(x*1600),int(y*900),int((x+w)*1600),int((y+h)*900));d.rounded_rectangle(box,radius=20,fill=bg)
 d.multiline_text((box[0]+24,box[1]+25),text,font=fonts.load_font('黑体',40 if name=='title' else 30),fill='white' if name=='title' else '#132c4a',spacing=28)
png=out/'示例图片页.png';im.save(png)
slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.shapes.add_picture(str(png),0,0,width=prs.slide_width,height=prs.slide_height)
prs.save(out/'示例课件.pptx')
print(out)
