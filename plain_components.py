"""Rebuild an ordinary slide from content; never default to styled crops."""
import math,re
from PIL import Image, ImageDraw, ImageOps
import video_fonts as fonts


def spec(item):
    value=item.get('plain')
    if value is not None:
        return value
    if isinstance(item.get('text'),str) and item['text'].strip():
        return {'kind':'text','text':item['text']}
    raise ValueError('普通版尚未重做：请填写原文、重建表格，或明确选择独立照片/标志；不能直接复用成品排版图块')


def rect(value):
    if (not isinstance(value,list) or len(value)!=4 or
        any(type(x) not in (int,float) or not math.isfinite(x) for x in value)):
        raise ValueError('普通版子组件位置无效')
    x,y,w,h=value
    if min(x,y)<0 or min(w,h)<=0 or x+w>1.001 or y+h>1.001:
        raise ValueError('普通版子组件须完整位于容器内')


def validate(value,depth=0):
    if not isinstance(value,dict):raise ValueError('请填写普通版内容')
    kind=value.get('kind')
    if kind=='text':
        if not isinstance(value.get('text'),str) or not value['text'].strip() or len(value['text'])>4000:
            raise ValueError('普通版文字须为1–4000字的原文')
    elif kind=='table':
        rows=value.get('rows')
        if not isinstance(rows,list) or not 1<=len(rows)<=30:raise ValueError('普通表格须有1–30行')
        if not isinstance(rows[0],list) or not 1<=len(rows[0])<=12:raise ValueError('普通表格须有1–12列')
        if any(not isinstance(row,list) or len(row)!=len(rows[0]) or
               any(not isinstance(cell,str) or len(cell)>300 for cell in row) for row in rows):
            raise ValueError('普通表格每行列数必须一致，单元格须为文字且不超过300字')
        widths=value.get('columns')
        if widths is not None and (not isinstance(widths,list) or len(widths)!=len(rows[0]) or
                any(type(x) not in (int,float) or not math.isfinite(x) or x<=0 for x in widths)):
            raise ValueError('表格列宽须与列数一致且为正数')
    elif kind=='image':
        if value.get('assetRole') not in ('photo','logo'):
            raise ValueError('仅独立照片或标志可沿用原图；文字、表格、图表和整块设计须重建普通版')
        rect(value.get('crop',[0,0,1,1]))
    elif kind=='group':
        parts=value.get('parts')
        if depth or not isinstance(parts,list) or not 1<=len(parts)<=20:
            raise ValueError('普通版组合须有1–20个子组件，不支持嵌套组合')
        for part in parts:
            if not isinstance(part,dict):raise ValueError('普通版子组件无效')
            rect(part.get('box'));validate(part,depth+1)
    else:raise ValueError('请选择文字、表格、独立照片/标志或已重建组合')
    if not isinstance(value.get('title',''),str) or len(value.get('title',''))>200:
        raise ValueError('普通版小标题最多200字')
    if value.get('role','body') not in ('title','body','caption'):
        raise ValueError('普通版文字层级无效')


def has_rebuilt_content(value):
    return value['kind'] in ('text','table') or (value['kind']=='group' and
            any(has_rebuilt_content(part) for part in value['parts']))


def wrap(draw,text,font,width):
    lines=[];closing='、，。；：！？）》】”’%'
    for paragraph in text.split('\n'):
        result=[];line=''
        tokens=re.findall(r'[A-Za-z0-9_./%+-]+|.',paragraph)
        for token in tokens:
            pieces=list(token) if draw.textlength(token,font=font)>width else [token]
            for char in pieces:
                if line and draw.textlength(line+char,font=font)>width:
                    if char[0] in closing and len(line)>1:
                        result.append(line[:-1]);line=line[-1]+char
                    else:result.append(line);line=char
                else:line+=char
        result.append(line)
        # Avoid a final line containing only one or two CJK characters.
        if len(result)>1 and 0<len(result[-1].strip())<3 and all('\u3400'<=c<='\u9fff' or c in closing for c in result[-1].strip()):
            amount=4-len(result[-1].strip())
            if len(result[-2])>amount+2:result[-1]=result[-2][-amount:]+result[-1];result[-2]=result[-2][:-amount]
        lines.extend(result)
    return lines


def text_layout(text,size,role='body'):
    w,h=size;draw=ImageDraw.Draw(Image.new('RGB',(1,1)))
    cap={'title':34,'body':28,'caption':17}[role]
    for n in range(cap,7,-1):
        font=fonts.load_font('黑体',n)
        lines=wrap(draw,text,font,max(1,w-8))
        leading=n+4
        if len(lines)*leading<=h-8 and all(draw.textlength(line,font=font)<=w-8 for line in lines):
            return font,lines,leading
    raise ValueError('普通版原文放不下，请扩大区域或重新分配版面；不能删减内容')


def table_layout(value,size):
    w,h=size;rows=value['rows'];widths=value.get('columns',[1]*len(rows[0]));total=sum(widths)
    xs=[0]
    for x in widths:xs.append(xs[-1]+x/total*w)
    xs[-1]=w
    draw=ImageDraw.Draw(Image.new('RGB',(1,1)))
    for n in range(20,7,-1):
        font=fonts.load_font('黑体',n);leading=n+3
        lines=[[wrap(draw,cell,font,max(1,xs[j+1]-xs[j]-10)) for j,cell in enumerate(row)] for row in rows]
        heights=[max(len(cell) for cell in row)*leading+8 for row in lines]
        fits=all(draw.textlength(line,font=font)<=xs[j+1]-xs[j]-10
                 for row in lines for j,cell in enumerate(row) for line in cell)
        if sum(heights)<=h and fits:return font,lines,leading,xs,heights
    raise ValueError('普通版表格放不下，请扩大表格区域；不能省略行列或数据')


def asset(bitmap,value):
    x,y,w,h=value.get('crop',[0,0,1,1])
    return bitmap.crop((round(x*bitmap.width),round(y*bitmap.height),
                        round((x+w)*bitmap.width),round((y+h)*bitmap.height)))


def render(value,bitmap,size):
    w,h=size;tile=Image.new('RGB',size,'white');draw=ImageDraw.Draw(tile)
    kind=value['kind'];offset=0
    if value.get('title'):
        title_h=min(30,max(28,h//7))
        title=render({'kind':'text','text':value['title'],'role':'body'},bitmap,(w,title_h))
        tile.paste(title,(0,0));offset=title_h+3
    body_size=(w,max(1,h-offset))
    if kind=='text':
        font,lines,leading=text_layout(value['text'],body_size,value.get('role','body'))
        for i,line in enumerate(lines):draw.text((4,offset+4+i*leading),line,font=font,fill='#222222')
    elif kind=='table':
        font,lines,leading,xs,heights=table_layout(value,body_size);y=offset
        for row,rh in zip(lines,heights):
            for j,cell in enumerate(row):
                draw.rectangle((round(xs[j]),y,min(w-1,round(xs[j+1])),y+rh-1),outline='#555555',width=1)
                for k,line in enumerate(cell):draw.text((round(xs[j])+5,y+4+k*leading),line,font=font,fill='#222222')
            y+=rh
    elif kind=='image':
        im=ImageOps.contain(asset(bitmap,value),(max(1,w-8),max(1,h-offset-8)),Image.Resampling.LANCZOS)
        tile.paste(im,((w-im.width)//2,offset+(h-offset-im.height)//2))
    elif kind=='group':
        for part in value['parts']:
            x,y,pw,ph=part['box'];left,top=round(x*w),offset+round(y*(h-offset))
            part_size=(max(1,round(pw*w)),max(1,round(ph*(h-offset))))
            tile.paste(render(part,bitmap,part_size),(left,top))
    return tile


def add_ppt(slide,value,bitmap,left,top,width,height):
    """Keep plain text and tables editable in the accompanying ordinary PPT."""
    import io
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR
    from pptx.oxml.xmlchemy import OxmlElement
    if value.get('title'):
        pixel_height=round(height/914400*80)
        th=round((min(30,max(28,pixel_height//7))+3)/80*914400)
        add_ppt(slide,{'kind':'text','text':value['title']},bitmap,left,top,width,th)
        top+=th;height-=th
    kind=value['kind']
    if kind=='group':
        for part in value['parts']:
            x,y,w,h=part['box'];add_ppt(slide,part,bitmap,int(left+x*width),int(top+y*height),int(w*width),int(h*height))
    elif kind=='image':
        image=asset(bitmap,value);factor=min(width/image.width,height/image.height)
        iw,ih=round(image.width*factor),round(image.height*factor)
        data=io.BytesIO();image.save(data,format='PNG');data.seek(0)
        slide.shapes.add_picture(data,int(left+(width-iw)/2),int(top+(height-ih)/2),iw,ih)
    elif kind=='text':
        shape=slide.shapes.add_textbox(left,top,width,height);tf=shape.text_frame
        tf.word_wrap=True;tf.margin_left=tf.margin_right=Pt(3);tf.margin_top=tf.margin_bottom=Pt(3)
        # The raster canvas uses 80px per inch, PPT uses 72pt per inch.
        px=(max(1,round(width/914400*80)),max(1,round(height/914400*80)))
        font,_,_=text_layout(value['text'],px,value.get('role','body'))
        tf.text=value['text']
        for p in tf.paragraphs:
            p.font.name='Noto Sans CJK SC';p.font.size=Pt(font.size*.9);p.font.color.rgb=RGBColor(34,34,34)
            p.space_before=p.space_after=Pt(0)
    elif kind=='table':
        rows=value['rows'];shape=slide.shapes.add_table(len(rows),len(rows[0]),left,top,width,height);table=shape.table
        table.first_row=False;table.horz_banding=False
        px=(max(1,round(width/914400*80)),max(1,round(height/914400*80)))
        font,_,_,xs,heights=table_layout(value,px)
        for j,col in enumerate(table.columns):col.width=int(width*(xs[j+1]-xs[j])/px[0])
        for i,row in enumerate(rows):
            table.rows[i].height=int(heights[i]/80*914400)
            for j,text in enumerate(row):
                cell=table.cell(i,j);cell.text=text;cell.fill.solid();cell.fill.fore_color.rgb=RGBColor(255,255,255)
                cell.margin_left=cell.margin_right=Pt(3);cell.margin_top=cell.margin_bottom=Pt(2)
                cell.vertical_anchor=MSO_ANCHOR.TOP
                for p in cell.text_frame.paragraphs:
                    p.font.name='Noto Sans CJK SC';p.font.size=Pt(font.size*.9);p.font.color.rgb=RGBColor(34,34,34)
                tcpr=cell._tc.get_or_add_tcPr()
                for edge in ('L','R','T','B'):
                    line=OxmlElement('a:ln'+edge);line.set('w','6350')
                    fill=OxmlElement('a:solidFill');color=OxmlElement('a:srgbClr');color.set('val','555555');fill.append(color);line.append(fill);tcpr.append(line)
